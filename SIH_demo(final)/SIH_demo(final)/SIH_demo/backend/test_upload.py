import requests
from docx import Document
import os

# Create dummy document
doc = Document()
doc.add_paragraph("Registration Certificate")
doc.add_paragraph("Registration Number: 27AAAAA0000A1Z5")
doc.add_paragraph("1. Legal Name: ACME Corp")
doc.add_paragraph("2. Trade Name: ACME Trading")
doc.add_paragraph("3. Constitution of Business: Private Limited Company")
doc.add_paragraph("5. Address of Principal Place of Business: 123 Tech Park, Mumbai")
doc.add_paragraph("6. Type of Registration: Regular")
doc.add_paragraph("7. Date of Issue: 01-Jan-2023")

doc_path = "dummy_gst.docx"
doc.save(doc_path)

# Test upload
url = "http://localhost:8000/upload/gst"
with open(doc_path, "rb") as f:
    files = {"file": ("dummy_gst.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = requests.post(url, files=files)

print("Status Code:", response.status_code)
import json
print("Response JSON:")
print(json.dumps(response.json(), indent=2))

# Cleanup
os.remove(doc_path)
