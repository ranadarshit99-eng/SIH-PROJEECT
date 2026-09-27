import shutil
from pathlib import Path
from typing import Tuple

import pymupdf as fitz
from docx import Document
from fastapi import UploadFile
from openpyxl import load_workbook
from pptx import Presentation

UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)


def save_and_extract(file: UploadFile) -> Tuple[str, str, str]:
    """Save uploaded file and extract its text content.

    Returns:
        (raw_text, filename, extension)
    """
    filename = file.filename
    extension = filename.split(".")[-1].lower()
    file_path = UPLOAD_DIR / filename

    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    raw_text = _extract_text(file_path, extension)
    return raw_text, filename, extension


def _extract_text(file_path: Path, extension: str) -> str:
    """Extract plain text from a file based on its extension."""
    text = ""
    if extension == "pdf":
        document = fitz.open(file_path)
        for page in document:
            text += page.get_text()
    elif extension == "docx":
        document = Document(file_path)
        for paragraph in document.paragraphs:
            text += paragraph.text + "\n"
    elif extension == "xlsx":
        workbook = load_workbook(file_path)
        for sheet in workbook:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        text += str(cell.value) + " "
    elif extension == "pptx":
        presentation = Presentation(file_path)
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
    return text
